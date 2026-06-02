"""
Step 1 — Document parser with tiered parsing strategy.

Priority order:
  1. LlamaParse  (cloud, best quality — needs LLAMA_CLOUD_API_KEY env var)
  2. Unstructured (local, good quality — needs `pip install unstructured[docx]`)
  3. python-docx  (fallback, always available)

All parsers return the same dict:
  {full_text, sections, tables, parser_used}
"""

import sys, os, re, io
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))


# ── Invalid XML char cleaner (shared) ────────────────────────────────────────

_INVALID_XML = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F\uFFFE\uFFFF]')

def _clean(text: str) -> str:
    text = _INVALID_XML.sub("", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"(?im)^page \d+ of \d+$", "", text)
    text = re.sub(r"(?im)^confidential.*$", "", text)
    return text.strip()


# ── Parser 1: LlamaParse ──────────────────────────────────────────────────────

def _parse_llamaparse(file_bytes: bytes) -> dict | None:
    """
    Uses LlamaParse cloud API to convert .docx to clean markdown.
    Requires: pip install llama-parse
    Requires: LLAMA_CLOUD_API_KEY environment variable
    """
    api_key = os.environ.get("LLAMA_CLOUD_API_KEY", "")
    if not api_key:
        return None
    try:
        from llama_parse import LlamaParse
        import tempfile

        parser = LlamaParse(
            api_key=api_key,
            result_type="markdown",
            verbose=False,
            language="en",
        )
        # LlamaParse needs a file path
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            documents = parser.load_data(tmp_path)
            markdown  = "\n\n".join(doc.text for doc in documents)
        finally:
            os.unlink(tmp_path)

        sections = _sections_from_markdown(markdown)
        return {
            "full_text":   _clean(markdown),
            "sections":    sections,
            "tables":      _extract_tables_from_markdown(markdown),
            "parser_used": "LlamaParse",
        }
    except Exception as e:
        print(f"[parser] LlamaParse failed: {e}")
        return None


# ── Parser 2: Unstructured ────────────────────────────────────────────────────

def _parse_unstructured(file_bytes: bytes) -> dict | None:
    """
    Uses Unstructured library for local high-quality docx parsing.
    Requires: pip install "unstructured[docx]"
    """
    try:
        from unstructured.partition.docx import partition_docx
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            elements = partition_docx(filename=tmp_path)
        finally:
            os.unlink(tmp_path)

        # Convert elements to markdown-like text
        parts   = []
        sections = []
        tables  = []
        current_heading = "Introduction"
        current_paras   = []

        for el in elements:
            etype = type(el).__name__
            text  = str(el).strip()
            if not text:
                continue

            if etype in ("Title", "Header"):
                if current_paras:
                    sections.append({"heading": current_heading, "content": " ".join(current_paras)})
                current_heading = text
                current_paras   = []
                parts.append(f"\n## {text}\n")
            elif etype == "Table":
                tables.append(text)
                parts.append(f"\n{text}\n")
            else:
                current_paras.append(text)
                parts.append(text)

        if current_paras:
            sections.append({"heading": current_heading, "content": " ".join(current_paras)})

        full_text = _clean("\n".join(parts))
        return {
            "full_text":   full_text,
            "sections":    sections,
            "tables":      tables,
            "parser_used": "Unstructured",
        }
    except ImportError:
        return None
    except Exception as e:
        print(f"[parser] Unstructured failed: {e}")
        return None


# ── Parser 3: python-docx (fallback) ─────────────────────────────────────────

def _parse_docx_fallback(file_bytes: bytes) -> dict:
    from docx import Document

    doc = None
    for attempt in [file_bytes, _sanitise_bytes(file_bytes)]:
        try:
            doc = Document(io.BytesIO(attempt))
            break
        except Exception:
            continue

    if doc is None:
        raise ValueError(
            "Could not parse the .docx file. "
            "The file may be corrupted or password-protected."
        )

    sections    = []
    tables_text = []
    current_heading = "Introduction"
    current_paras   = []

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            tables_text.append("\n".join(rows))

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower() if para.style else ""
        is_heading = (
            "heading" in style_name
            or (para.runs and any(r.bold for r in para.runs))
        )
        if is_heading and len(text) < 120:
            if current_paras:
                sections.append({"heading": current_heading, "content": " ".join(current_paras)})
            current_heading = text
            current_paras   = []
        else:
            current_paras.append(text)

    if current_paras:
        sections.append({"heading": current_heading, "content": " ".join(current_paras)})

    parts = [f"## {s['heading']}\n{s['content']}" for s in sections]
    for i, t in enumerate(tables_text):
        parts.append(f"## Table {i+1}\n{t}")

    return {
        "full_text":   _clean("\n\n".join(parts)),
        "sections":    sections,
        "tables":      tables_text,
        "parser_used": "python-docx",
    }


def _sanitise_bytes(data: bytes) -> bytes:
    try:
        text = data.decode("utf-8", errors="replace")
        text = _INVALID_XML.sub("", text)
        return text.encode("utf-8")
    except Exception:
        return data


# ── Markdown helpers ──────────────────────────────────────────────────────────

def _sections_from_markdown(md: str) -> list[dict]:
    sections = []
    current_heading = "Introduction"
    current_lines   = []
    for line in md.splitlines():
        if line.startswith("#"):
            if current_lines:
                sections.append({"heading": current_heading, "content": " ".join(current_lines)})
            current_heading = line.lstrip("#").strip()
            current_lines   = []
        elif line.strip():
            current_lines.append(line.strip())
    if current_lines:
        sections.append({"heading": current_heading, "content": " ".join(current_lines)})
    return sections


def _extract_tables_from_markdown(md: str) -> list[str]:
    tables = []
    in_table = False
    rows = []
    for line in md.splitlines():
        if "|" in line:
            in_table = True
            rows.append(line.strip())
        elif in_table:
            if rows:
                tables.append("\n".join(rows))
            rows = []
            in_table = False
    if rows:
        tables.append("\n".join(rows))
    return tables


# ── Main entry point ──────────────────────────────────────────────────────────

def parse_docx(file) -> dict:
    """
    Parse a .docx file using the best available parser.
    Falls back gracefully: LlamaParse → Unstructured → python-docx
    """
    raw_bytes = file.read() if hasattr(file, "read") else file

    for parser_fn in [_parse_llamaparse, _parse_unstructured]:
        result = parser_fn(raw_bytes)
        if result:
            print(f"[parser] Used: {result['parser_used']}")
            return result

    # Guaranteed fallback
    result = _parse_docx_fallback(raw_bytes)
    print(f"[parser] Used: {result['parser_used']}")
    return result